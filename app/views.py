from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from .forms import UserSignupForm, LawyerSignupForm, LoginForm
from .models import CustomUser, LawyerProfile, Hire, Message


# -------- User Signup --------
def user_signup(request):
    form = UserSignupForm(request.POST or None)
    if form.is_valid():
        form.save()
        return redirect('login')
    return render(request, 'signup_user.html', {'form': form})


# -------- Lawyer Signup --------
def lawyer_signup(request):
    form = LawyerSignupForm(request.POST or None, request.FILES or None)
    
    if request.method == 'POST':
        print("Received POST request")
        if form.is_valid():
            print("Form is valid")
            user = form.save()
            print("User created:", user)
            return redirect('login')
        else:
            print("Form is NOT valid")
            print(form.errors) 

    return render(request, 'signup_lawyer.html', {'form': form})


# -------- Login View --------
def login_view(request):
    form = LoginForm(request.POST or None)
    error = ""

    if request.method == "POST" and form.is_valid():
        email = form.cleaned_data['email']
        password = form.cleaned_data['password']
        user = authenticate(request, email=email, password=password)

        if user is not None:
            if user.role == 'lawyer':
                try:
                    if not user.lawyerprofile.is_approved:
                        error = "Your account is awaiting admin approval."
                        return render(request, 'login.html', {'form': form, 'error': error})
                except LawyerProfile.DoesNotExist:
                    error = "Lawyer profile not found. Contact admin."
                    return render(request, 'login.html', {'form': form, 'error': error})

            login(request, user)

            
            if user.role == 'admin':
                return redirect('admin_dashboard')
            elif user.role == 'lawyer':
                return redirect('lawyer_dashboard')
            elif user.role == 'user':
                return redirect('user_dashboard')
            else:
                error = "Invalid user role. Contact admin."
        else:
            error = "Invalid email or password."

    return render(request, 'login.html', {'form': form, 'error': error})


# -------- User Dashboard --------
@login_required
def user_dashboard(request):
    if request.user.role != 'user':
        return redirect('login')
    return render(request, 'user_dashboard.html', {
        'name': request.user.full_name,
    })


# -------- Lawyer Dashboard --------
@login_required
def lawyer_dashboard(request):
    if request.user.role != 'lawyer':
        return redirect('login')
    return render(request, 'lawyer_dashboard.html', {
        'name': request.user.full_name,
    })


# -------- Admin Dashboard --------
@login_required
def admin_dashboard(request):
    if request.user.role != 'admin':
        return redirect('login')
    
    # Get counts for dashboard cards
    pending_count = LawyerProfile.objects.filter(is_approved=False).count()
    approved_count = LawyerProfile.objects.filter(is_approved=True).count()
    total_users = CustomUser.objects.filter(role='user').count()
    
    return render(request, 'admin_dashboard.html', {
        'pending_count': pending_count,
        'approved_count': approved_count,
        'total_users': total_users,
    })

@login_required
def pending_approvals(request):
    if request.user.role != 'admin':
        return redirect('login')
    pending_lawyers = LawyerProfile.objects.filter(is_approved=False)
    return render(request, 'pending_approvals.html', {
        'pending_lawyers': pending_lawyers,
    })


@login_required
def manage_lawyers(request):
    if request.user.role != 'admin':
        return redirect('login')
    approved_lawyers = LawyerProfile.objects.filter(is_approved=True)
    return render(request, 'manage_lawyers.html', {
        'approved_lawyers': approved_lawyers,
    })


from django.core.mail import send_mail
from django.conf import settings
from django.contrib import messages


@login_required
def deny_lawyer(request, lawyer_id):
    if request.user.role != 'admin':
        return redirect('login')
    try:
        profile = LawyerProfile.objects.get(id=lawyer_id)
        lawyer_name = profile.user.full_name
        lawyer_email = profile.user.email
        
        
        try:
            send_mail(
                subject='Application Denied - AILA Legal Platform',
                message=f'''Dear {lawyer_name},

We regret to inform you that your application to join AILA Legal Platform has been denied after careful review.

This decision may be due to:
- Incomplete documentation
- Verification requirements not met
- Platform eligibility criteria

You may reapply after addressing these concerns.

Best regards,
AILA Admin Team''',
                from_email=settings.EMAIL_HOST_USER,
                recipient_list=[lawyer_email],
                fail_silently=False,
            )
        except Exception as e:
            messages.warning(request, f"Lawyer denied but email notification failed: {str(e)}")
        
        # Delete the profile and user
        user = profile.user
        profile.delete()
        user.delete()
        
        messages.success(request, f"Lawyer {lawyer_name} has been denied and removed.")
    except LawyerProfile.DoesNotExist:
        messages.error(request, "Lawyer profile not found.")
    
    return redirect('pending_approvals')

@login_required
def send_lawyer_report(request, lawyer_id):
    if request.user.role != 'admin':
        return redirect('login')
    
    if request.method == 'POST':
        try:
            profile = LawyerProfile.objects.get(id=lawyer_id)
            report_subject = request.POST.get('report_subject', '').strip()
            report_message = request.POST.get('report_message', '').strip()
            
            if report_subject and report_message:
                try:
                    send_mail(
                        subject=f'AILA Platform Report: {report_subject}',
                        message=f'''Dear {profile.user.full_name},

{report_message}

This is an official communication from AILA Administration.

Best regards,
AILA Admin Team''',
                        from_email=settings.EMAIL_HOST_USER,
                        recipient_list=[profile.user.email],
                        fail_silently=False,
                    )
                    messages.success(request, f"Report sent successfully to {profile.user.full_name}")
                except Exception as e:
                    messages.error(request, f"Failed to send report: {str(e)}")
            else:
                messages.error(request, "Please provide both subject and message.")
        except LawyerProfile.DoesNotExist:
            messages.error(request, "Lawyer profile not found.")
    
    return redirect('manage_lawyers')


@login_required
def remove_lawyer(request, lawyer_id):
    if request.user.role != 'admin':
        return redirect('login')
    
    if request.method == 'POST':
        try:
            profile = LawyerProfile.objects.get(id=lawyer_id)
            lawyer_name = profile.user.full_name
            lawyer_email = profile.user.email
            removal_reason = request.POST.get('removal_reason', '').strip()
            
            
            try:
                send_mail(
                    subject='Account Removed - AILA Legal Platform',
                    message=f'''Dear {lawyer_name},

Your account on AILA Legal Platform has been removed by administration.

Reason: {removal_reason if removal_reason else 'Administrative decision'}

If you believe this is an error, please contact our support team.

Best regards,
AILA Admin Team''',
                    from_email=settings.EMAIL_HOST_USER,
                    recipient_list=[lawyer_email],
                    fail_silently=False,
                )
            except Exception as e:
                messages.warning(request, f"Lawyer removed but email notification failed: {str(e)}")
            
            
            user = profile.user
            profile.delete()
            user.delete()
            
            messages.success(request, f"Lawyer {lawyer_name} has been removed successfully.")
        except LawyerProfile.DoesNotExist:
            messages.error(request, "Lawyer profile not found.")
    
    return redirect('manage_lawyers')

login_required
def approve_lawyer(request, lawyer_id):
    if request.user.role != 'admin':
        return redirect('login')
    
    try:
        profile = LawyerProfile.objects.get(id=lawyer_id, is_approved=False)
        profile.is_approved = True
        profile.save()
        
        try:
            send_mail(
                subject='Application Approved - AILA Legal Platform',
                message=f'''Dear {profile.user.full_name},

Congratulations! Your application to join AILA Legal Platform has been approved.

You can now:
- Log in to your lawyer dashboard
- Set your availability
- Receive client inquiries
- Manage your cases

Welcome to the AILA family!

Best regards,
AILA Admin Team''',
                from_email=settings.EMAIL_HOST_USER,
                recipient_list=[profile.user.email],
                fail_silently=False,
            )
            messages.success(request, f"Lawyer {profile.user.full_name} has been approved and notified via email.")
        except Exception as e:
            messages.warning(request, f"Lawyer approved but email notification failed: {str(e)}")
            
    except LawyerProfile.DoesNotExist:
        messages.error(request, "Lawyer profile not found or already approved.")
    
    return redirect('pending_approvals')


# -------- Logout --------
def logout_view(request):
    logout(request)
    return redirect('login')

@login_required
def find_lawyer(request):
    if request.user.role != 'user':
        return redirect('login')
    
    from django.db.models import Avg, Count
    
    lawyers = LawyerProfile.objects.filter(is_approved=True).annotate(
        average_rating=Avg('user__lawyer_reviews__rating'),
        review_count=Count('user__lawyer_reviews')
    )
    
    for lawyer in lawyers:
        if lawyer.average_rating is None:
            lawyer.average_rating = 0
        else:
            lawyer.average_rating = round(lawyer.average_rating, 1)
        if lawyer.review_count is None:
            lawyer.review_count = 0
            
    return render(request, 'find_lawyer.html', {'lawyers': lawyers})

@login_required
def hire_lawyer(request, lawyer_id):
    if request.user.role != 'user':
        return redirect('login')

    lawyer = CustomUser.objects.get(id=lawyer_id)
    if not Hire.objects.filter(user=request.user, lawyer=lawyer).exists():
        Hire.objects.create(user=request.user, lawyer=lawyer)
    return redirect('contacted_lawyers')

@login_required
def contacted_lawyers(request):
    if request.user.role != 'user':
        return redirect('login')
    hires = Hire.objects.filter(user=request.user)
    return render(request, 'contacted_lawyers.html', {'hires': hires})

from channels.layers import get_channel_layer
from asgiref.sync import async_to_sync
from django.shortcuts import get_object_or_404, render, redirect
from django.db.models import Q
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse

from .models import CustomUser, Message
@login_required
def chat_view(request, receiver_id):
    receiver = get_object_or_404(CustomUser, id=receiver_id)

    messages = Message.objects.filter(
        sender__in=[request.user, receiver],
        receiver__in=[request.user, receiver]
    ).exclude(
        Q(sender=request.user, deleted_by_sender=True) |
        Q(receiver=request.user, deleted_by_receiver=True)
    ).order_by('timestamp')

    if request.method == 'POST':
        content = request.POST.get('content', '')
        uploaded_file = request.FILES.get('file')

        if content or uploaded_file:
            msg = Message.objects.create(
                sender=request.user,
                receiver=receiver,
                content=content,
                file=uploaded_file
            )

            channel_layer = get_channel_layer()
            if receiver.role == 'user':
                async_to_sync(channel_layer.group_send)(
                    f"user_list_{receiver.id}",
                    {
                        'type': 'user_list_update',
                        'data': {
                            'action': 'new_message',
                            'sender_id': request.user.id,
                            'sender_name': request.user.full_name,
                            'sender_email': request.user.email,
                            'last_message': msg.content[:50] + '...' if len(msg.content) > 50 else msg.content,
                            'timestamp': msg.timestamp.strftime("%H:%M"),
                        }
                    }
                )

            async_to_sync(channel_layer.group_send)(
                f"chat_{receiver.id}_{request.user.id}",
                {
                    'type': 'chat_message',
                    'message': msg.content,
                    'sender': request.user.full_name,
                    'timestamp': msg.timestamp.strftime('%b %d, %H:%M'),
                    'file_url': msg.file.url if msg.file else '',
                    'is_self': True 
                }
            )

            async_to_sync(channel_layer.group_send)(
                f"chat_{request.user.id}_{receiver.id}",
                {
                    'type': 'chat_message',
                    'message': msg.content,
                    'sender': request.user.full_name,
                    'timestamp': msg.timestamp.strftime('%b %d, %H:%M'),
                    'file_url': msg.file.url if msg.file else '',
                    'is_self': False  
                }
            )

        return HttpResponse(status=200)

    return render(request, 'chat.html', {
        'receiver': receiver,
        'messages': messages
    })

@login_required
def messaged_users(request):
    if request.user.role != 'lawyer':
        return redirect('login')

    from django.db.models import Max, Subquery, OuterRef, Count
    
    last_message_subquery = Message.objects.filter(
        Q(sender=OuterRef('pk'), receiver=request.user) | 
        Q(sender=request.user, receiver=OuterRef('pk'))
    ).exclude(
        Q(sender=request.user, deleted_by_sender=True) |
        Q(receiver=request.user, deleted_by_receiver=True)
    ).order_by('-timestamp')

    users_with_messages = CustomUser.objects.filter(
        sent_messages__receiver=request.user,
        sent_messages__deleted_by_receiver=False
    ).annotate(
        last_message_time=Max('sent_messages__timestamp'),
        last_message_content=Subquery(last_message_subquery.values('content')[:1]),
        last_message_timestamp=Subquery(last_message_subquery.values('timestamp')[:1]),
        unread_count=Count('sent_messages', filter=Q(
            sent_messages__receiver=request.user,
            sent_messages__is_read=False,
            sent_messages__deleted_by_receiver=False
        ))
    ).distinct().order_by('-last_message_time')

    
    messages = []
    selected_user = None

    user_id = request.GET.get('user_id')
    if user_id:
        try:
            selected_user = CustomUser.objects.get(id=user_id)
            messages = Message.objects.filter(
                sender__in=[request.user, selected_user],
                receiver__in=[request.user, selected_user]
            ).exclude(
                Q(sender=request.user, deleted_by_sender=True) |
                Q(receiver=request.user, deleted_by_receiver=True)
            ).order_by('timestamp')
            
            Message.objects.filter(
                sender=selected_user,
                receiver=request.user,
                is_read=False
            ).update(is_read=True)
            
        except CustomUser.DoesNotExist:
            selected_user = None

    return render(request, 'messaged_users.html', {
        'users': users_with_messages,  
        'messages': messages,
        'selected_user': selected_user
    })

from django.shortcuts import redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from .models import LawyerReview
from django.contrib.auth import get_user_model
from django.shortcuts import get_object_or_404

User = get_user_model()

@login_required
def submit_review(request, lawyer_id):
    if request.method == 'POST':
        review_text = request.POST.get('review', '')
        rating = int(request.POST.get('rating', 0))
        lawyer = get_object_or_404(User, pk=lawyer_id)

        if review_text and 1 <= rating <= 5:
            LawyerReview.objects.create(
                lawyer=lawyer,
                reviewer=request.user,
                review=review_text,
                rating=rating
            )

    return redirect('contacted_lawyers')  


@login_required
def lawyer_profile(request, lawyer_id):
    from django.db.models import Avg, Count
    
    lawyer = get_object_or_404(CustomUser, id=lawyer_id, role='lawyer')
    profile = get_object_or_404(LawyerProfile, user=lawyer)
    reviews = LawyerReview.objects.filter(lawyer=lawyer).order_by('-created_at')
    

    rating_stats = LawyerReview.objects.filter(lawyer=lawyer).aggregate(
        average_rating=Avg('rating'),
        review_count=Count('id')
    )
    
    average_rating = rating_stats['average_rating'] or 0
    review_count = rating_stats['review_count'] or 0
    
    return render(request, 'lawyer_profile.html', {
        'lawyer': lawyer,
        'profile': profile,
        'reviews': reviews,
        'average_rating': round(average_rating, 1) if average_rating else 0,
        'review_count': review_count
    })

def home(request):
    return render(request, 'home.html')

def about(request):
    return render(request, 'about.html')

def nda(request):
    return render(request, 'nda.html')

def poa(request):
    return render(request, 'poa.html')

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from .models import Question, Answer
from .forms import QuestionForm
from django.contrib import messages
from .utils import send_answer_notification 
def user_qa(request):
    """Public Q&A - No login required"""
    
    if request.method == 'POST':
        form = QuestionForm(request.POST)
        if form.is_valid():
            question = form.save(commit=False)
            question.user = request.user if request.user.is_authenticated else None
            question.save()
            
            messages.success(
                request, 
                f"Question submitted successfully! You will be emailed all responses on {question.email}"
            )
            return redirect('user_qa')
        else:
            messages.error(request, "Please correct the errors below.")
    else:
        form = QuestionForm()
    answered_questions = Question.objects.filter(answers__isnull=False).distinct().order_by('-created_at')

    return render(request, 'user_qa.html', {
        'form': form,
        'questions': answered_questions,
    })

from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from .models import Question, Answer 
from .utils import send_answer_notification  
from django.contrib import messages
from .utils import send_answer_notification  
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import Question, Answer
from .forms import QuestionForm
from .utils import send_answer_notification 

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from .models import Question, Answer
from .forms import QuestionForm
from django.contrib import messages
from .utils import send_answer_notification

def user_qa(request):
    """Public Q&A - No login required"""
    
    if request.method == 'POST':
        form = QuestionForm(request.POST)
        if form.is_valid():
            question = form.save(commit=False)
            question.user = None  
            question.save()
            
            messages.success(
                request, 
                f"Question submitted successfully! You will be emailed all responses on {question.email}"
            )
            return redirect('user_qa')
        else:
            messages.error(request, "Please correct the errors below.")
    else:
        form = QuestionForm()

    answered_questions = Question.objects.filter(answers__isnull=False).distinct().order_by('-created_at')

    return render(request, 'user_qa.html', {
        'form': form,
        'questions': answered_questions,
    })

@login_required
def lawyer_qa(request):
    """Lawyer Q&A - Requires login for lawyers only"""
   
    if request.user.role != 'lawyer':
        return redirect('login')

  
    if request.method == 'POST':
        question_id = request.POST.get('question_id')
        answer_text = request.POST.get('answer_text', '').strip()

        if question_id and answer_text:
            try:
                question = get_object_or_404(Question, id=question_id)

              
                existing_answer = Answer.objects.filter(
                    question=question,
                    lawyer=request.user
                ).first()

                if not existing_answer:
                    answer = Answer.objects.create(
                        question=question,
                        lawyer=request.user,
                        answer_text=answer_text
                    )
                    
                   
                    try:
                        if send_answer_notification(question, answer):
                            messages.success(request, "Answer submitted and email notification sent successfully!")
                        else:
                            messages.warning(request, "Answer submitted but email notification failed to send.")
                    except Exception as e:
                        messages.warning(request, f"Answer submitted but email notification failed: {str(e)}")
                else:
                    messages.info(request, "You have already answered this question.")
                    
            except Exception as e:
                messages.error(request, f"Error submitting answer: {str(e)}")
        else:
            messages.error(request, "Please provide both question ID and answer text.")

    unanswered_questions = Question.objects.filter(answers__isnull=True).order_by('created_at')

    answered_by_lawyer = Question.objects.filter(answers__lawyer=request.user).order_by('-answers__created_at')

    return render(request, 'lawyer_qa.html', {
        'unanswered_questions': unanswered_questions,
        'answered_by_lawyer': answered_by_lawyer
    })

@login_required
def edit_answer(request, answer_id):
    if request.user.role != 'lawyer':
        return redirect('login')
    
    try:
        answer = get_object_or_404(Answer, id=answer_id, lawyer=request.user)
        
        if request.method == 'POST':
            answer_text = request.POST.get('answer_text', '').strip()
            if answer_text:
                answer.answer_text = answer_text
                answer.save()
                
               
                try:
                    if send_answer_notification(answer.question, answer):
                        messages.success(request, "Answer updated and email notification sent!")
                    else:
                        messages.warning(request, "Answer updated but email notification failed to send.")
                except Exception as e:
                    messages.warning(request, f"Answer updated but email notification failed: {str(e)}")
            else:
                messages.error(request, "Answer text cannot be empty.")
    except Exception as e:
        messages.error(request, f"Error updating answer: {str(e)}")
    
    return redirect('lawyer_qa')

@login_required
def delete_answer(request, answer_id):
    if request.user.role != 'lawyer':
        return redirect('login')
    
    answer = get_object_or_404(Answer, id=answer_id, lawyer=request.user)
    answer.delete()
    
    return redirect('lawyer_qa')

from .models import Blog


def blog_list(request):
    query = request.GET.get('q', '')
    blogs = Blog.objects.all().order_by('-created_at')

    if query:
        blogs = blogs.filter(title__icontains=query) | blogs.filter(content__icontains=query)

    return render(request, 'blog.html', {
        'blogs': blogs,
        'query': query
    })


@login_required
def lawyer_blog(request):
    if request.user.role != 'lawyer':
        return redirect('login')

    if request.method == 'POST':
        title = request.POST.get('title', '').strip()
        content = request.POST.get('content', '').strip()
        image = request.FILES.get('image')

        if title and content:
            Blog.objects.create(lawyer=request.user, title=title, content=content, image=image)
            return redirect('lawyer_blog')

    blogs = Blog.objects.filter(lawyer=request.user).order_by('-created_at')
    return render(request, 'lawyer_blog.html', {
        'blogs': blogs
    })
@login_required
def edit_blog(request, blog_id):
    if request.user.role != 'lawyer':
        return redirect('login')
    
    try:
        blog = Blog.objects.get(id=blog_id, lawyer=request.user)
    except Blog.DoesNotExist:
        return redirect('lawyer_blog')
    
    if request.method == 'POST':
        title = request.POST.get('title', '').strip()
        content = request.POST.get('content', '').strip()
        image = request.FILES.get('image')
        
        if title and content:
            blog.title = title
            blog.content = content
            if image:
                blog.image = image
            blog.save()
            return redirect('lawyer_blog')
    
    return render(request, 'edit_blog.html', {'blog': blog})

@login_required
def delete_blog(request, blog_id):
    if request.user.role != 'lawyer':
        return redirect('login')
    
    try:
        blog = Blog.objects.get(id=blog_id, lawyer=request.user)
        blog.delete()
    except Blog.DoesNotExist:
        pass
    
    return redirect('lawyer_blog')

from django.shortcuts import render, redirect, get_object_or_404
from .models import Case 
from django.contrib import messages

def lawyer_case(request):
    cases = Case.objects.all()
    return render(request, 'lawyer_case.html', {'cases': cases})

from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from .models import Case
from django.contrib.auth import get_user_model

User = get_user_model()

@login_required
def create_case(request):
    if request.method == 'POST':
        email = request.POST['email']
        case_name = request.POST['case_name']
        description = request.POST['description']
        client_name = request.POST['client_name']  
        user = get_object_or_404(User, email=email)

        case = Case.objects.create(
            user=user,
            lawyer=request.user,
            case_name=case_name,
            description=description,
            client_name=client_name,  
            status='Reviewing'
        )
        return redirect('lawyer_case')
def update_case_status(request, case_id):
    case = get_object_or_404(Case, id=case_id)
    if request.method == 'POST':
        case.status = request.POST['status']
        case.save()
    return redirect('lawyer_case')
@login_required
def upload_case_document(request, case_id):
    case = get_object_or_404(Case, id=case_id)
    
    if request.method == 'POST':
        document = request.FILES.get('document')
        document_name = request.POST.get('document_name')
        
        if document and document_name:
            from .models import CaseDocument
            CaseDocument.objects.create(
                case=case,
                document=document,
                document_name=document_name,
                uploaded_by=request.user
            )
    
    return redirect('lawyer_case')

@login_required
def delete_case_document(request, document_id):
    from .models import CaseDocument
    document = get_object_or_404(CaseDocument, id=document_id)
    document.delete()
    return redirect('lawyer_case')

from django.contrib import messages

@login_required
def delete_case(request, case_id):
    case = get_object_or_404(Case, id=case_id)
    
    if request.method == 'POST':
        print(f"User role: {request.user.role}")
        print(f"Case lawyer: {case.lawyer}")
        print(f"Request user: {request.user}")
        print(f"Are they equal? {case.lawyer == request.user}")
        if request.user.role == 'lawyer':
            case_name = case.case_name
            case.delete()
            messages.success(request, f'Case "{case_name}" has been successfully deleted.')
            return redirect('lawyer_case')
        elif request.user.role == 'user' and case.user == request.user:
            case_name = case.case_name
            case.delete()
            messages.success(request, f'Case "{case_name}" has been successfully deleted.')
            return redirect('user_cases')
        else:
            messages.error(request, f'Permission denied. User role: {request.user.role}')
    
    return redirect('lawyer_case')

from django.contrib.auth.decorators import login_required
from django.shortcuts import render
from .models import Case

@login_required
def user_case_view(request):
    user = request.user
    cases = Case.objects.filter(user=user).order_by('-created_at')
    return render(request, 'user_case.html', {'cases': cases})

from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.http import JsonResponse
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth import get_user_model
from django.db.models import Q
from .models import AvailableSlot, Appointment, CustomUser
from .forms import SlotForm, AppointmentForm
from asgiref.sync import async_to_sync
from channels.layers import get_channel_layer

User = get_user_model()
@login_required
def set_appointment(request):
    if request.user.role != 'lawyer':
        return redirect('login')

    if request.method == 'POST':
        if 'delete_slot' in request.POST:
            slot_id = request.POST.get('delete_slot')
            try:
                slot = AvailableSlot.objects.get(id=slot_id, lawyer=request.user)
                slot_data = {
                    "id": slot.id,
                    "start_time": slot.start_time.isoformat(),
                    "end_time": slot.end_time.isoformat(),
                }
                
                slot.delete()
                messages.success(request, 'Slot deleted successfully!')
                channel_layer = get_channel_layer()
                async_to_sync(channel_layer.group_send)(
                    f"slots_{request.user.id}",
                    {
                        "type": "slot_update",
                        "data": {
                            "action": "slot_deleted",
                            "slot": slot_data
                        }
                    }
                )
                
            except AvailableSlot.DoesNotExist:
                messages.error(request, 'Slot not found!')
            return redirect('set_appointment')

        form = SlotForm(request.POST)
        if form.is_valid():
            start = form.cleaned_data['new_slot_start']
            end = form.cleaned_data['new_slot_end']

            if start >= end:
                messages.error(request, 'End time must be after start time!')
            elif start <= timezone.now():
                messages.error(request, 'Start time must be in the future!')
            else:
                overlapping_slots = AvailableSlot.objects.filter(
                    lawyer=request.user,
                    start_time__lt=end,
                    end_time__gt=start
                )

                if overlapping_slots.exists():
                    messages.error(request, 'This slot overlaps with an existing slot!')
                else:
                    new_slot = AvailableSlot.objects.create(
                        lawyer=request.user,
                        start_time=start,
                        end_time=end
                    )
                    messages.success(request, 'Slot added successfully!')

                   
                    channel_layer = get_channel_layer()
                    async_to_sync(channel_layer.group_send)(
                        f"slots_{request.user.id}",
                        {
                            "type": "slot_update",
                            "data": {
                                "action": "slot_added",
                                "slot": {
                                    "id": new_slot.id,
                                    "start_time": new_slot.start_time.isoformat(),
                                    "end_time": new_slot.end_time.isoformat(),
                                    "display": f"{new_slot.start_time.strftime('%b %d, %Y %H:%M')} – {new_slot.end_time.strftime('%H:%M')}"
                                }
                            }
                        }
                    )

            return redirect('set_appointment')
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"{field}: {error}")
    else:
        form = SlotForm()

    AvailableSlot.objects.filter(end_time__lt=timezone.now()).delete()

    slots = AvailableSlot.objects.filter(lawyer=request.user).order_by('start_time')

    slot_with_appointments = []
    for slot in slots:
        appointment = Appointment.objects.filter(slot=slot).first()
        slot_with_appointments.append({
            'slot': slot,
            'appointment': appointment,
        })

    return render(request, 'set_appointment.html', {
        'form': form,
        'slot_with_appointments': slot_with_appointments,
    })

@login_required
def book_appointment(request, lawyer_id):
    lawyer = get_object_or_404(CustomUser, id=lawyer_id, role='lawyer')

    booked_slot_ids = Appointment.objects.values_list('slot_id', flat=True)
    slots = AvailableSlot.objects.filter(
        lawyer=lawyer,
        end_time__gt=timezone.now()
    ).exclude(id__in=booked_slot_ids).order_by('start_time')

    if request.method == 'POST':
        form = AppointmentForm(request.POST)
        if form.is_valid():
            selected_time = form.cleaned_data['appointment_datetime']
            matching_slot = AvailableSlot.objects.filter(
                lawyer=lawyer,
                start_time__lte=selected_time,
                end_time__gte=selected_time,
                end_time__gt=timezone.now()
            ).exclude(id__in=booked_slot_ids).first()

            if matching_slot:
                appointment = Appointment.objects.create(slot=matching_slot, user=request.user)

                channel_layer = get_channel_layer()
                async_to_sync(channel_layer.group_send)(
                    f"slots_{lawyer.id}",
                    {
                        "type": "slot_update",
                        "data": {
                            "action": "slot_booked",
                            "slot_id": matching_slot.id,
                            "slot": {
                                "start_time": matching_slot.start_time.isoformat(),
                                "end_time": matching_slot.end_time.isoformat(),
                            },
                            "user_name": request.user.full_name,
                            "booked_at": appointment.booked_at.isoformat()
                        }
                    }
                )

                
                if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                    return JsonResponse({'success': True})

                messages.success(request, 'Appointment booked successfully!')
                return redirect('book_appointment', lawyer_id=lawyer.id)
            else:
                error_msg = 'Selected time is not available or already booked.'
                if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                    return JsonResponse({'success': False, 'error': error_msg})
                form.add_error(None, error_msg)
        else:
            if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                return JsonResponse({'success': False, 'error': 'Invalid form data'})
    else:
        form = AppointmentForm()

    return render(request, 'book_appointment.html', {
        'form': form,
        'slots': slots,
        'lawyer': lawyer,
    })

@login_required
def get_available_slots_api(request, lawyer_id):
    lawyer = get_object_or_404(CustomUser, id=lawyer_id, role='lawyer')
    booked_slot_ids = Appointment.objects.values_list('slot_id', flat=True)

    slots = AvailableSlot.objects.filter(
        lawyer=lawyer,
        end_time__gt=timezone.now()
    ).exclude(id__in=booked_slot_ids).order_by('start_time')

    slot_data = [
        {
            'id': slot.id,
            'start_time': slot.start_time.isoformat(),
            'end_time': slot.end_time.isoformat(),
            'display': f"{slot.start_time.strftime('%b %d, %Y %H:%M')} – {slot.end_time.strftime('%H:%M')}"
        }
        for slot in slots
    ]

    return JsonResponse({'slots': slot_data})

def docuement(request):
    return render(request, 'document.html')
def emp(request):
    return render(request, 'emp.html')
def doc(request):
    return render(request, 'document1.html')

def lease(request):
    return render(request, 'lease.html')


from .forms import LawyerProfileEditForm
from django.contrib.auth.decorators import login_required

@login_required
def edit_profile(request):
    if request.user.role != 'lawyer':
        return redirect('login')

    profile = LawyerProfile.objects.get(user=request.user)

    if request.method == 'POST':
        form = LawyerProfileEditForm(request.POST, request.FILES, instance=profile)
        if form.is_valid():
            form.save()
            return redirect('edit_profile')  
    else:
        form = LawyerProfileEditForm(instance=profile)

    return render(request, 'profile.html', {
        'form': form,
        'profile': profile
    })


from django.shortcuts import render, redirect
from .forms import FeedbackForm
from .models import Feedback
from django.contrib import messages
def submit_feedback(request):
    """Submit feedback - Lawyers must login, everyone else is anonymous"""
    if request.method == 'POST':
        form = FeedbackForm(request.POST, user=request.user if request.user.is_authenticated else None)
        
        if form.is_valid():
            feedback = form.save(commit=False)
            if (request.user.is_authenticated and 
                hasattr(request.user, 'role') and request.user.role == 'lawyer'):
                feedback.user = request.user
                feedback.anonymous_name = ''
                feedback.anonymous_email = ''
            else:
                feedback.user = None
                feedback.anonymous_name = 'Anonymous Visitor'
                feedback.anonymous_email = ''
            
            feedback.save()
            messages.success(request, 'Thank you for your feedback!')
            
            return redirect('submit_feedback')
        else:
            messages.error(request, 'Please correct the errors below.')
    else:
        form = FeedbackForm(user=request.user if request.user.is_authenticated else None)
    is_lawyer = (request.user.is_authenticated and 
                hasattr(request.user, 'role') and request.user.role == 'lawyer')
    
    return render(request, 'feedback.html', {
        'form': form,
        'is_lawyer': is_lawyer,
        'user_authenticated': request.user.is_authenticated
    })

@login_required
def view_feedback(request):
    """Admin view to see all feedback - lawyers vs anonymous"""
    if request.user.role != 'admin':
        return redirect('login')
    
    user_type = request.GET.get('type', 'all')  

    if user_type == 'lawyer':
        feedbacks = Feedback.objects.filter(user_type='lawyer').order_by('-submitted_at')
        filter_title = "Lawyer Feedback"
    elif user_type == 'visitor':
        feedbacks = Feedback.objects.filter(user_type='visitor').order_by('-submitted_at')
        filter_title = "Anonymous Feedback (Visitors & Users)"
    elif user_type == 'anonymous':
        
        feedbacks = Feedback.objects.filter(user__isnull=True).order_by('-submitted_at')
        filter_title = "Anonymous Feedback"
    elif user_type == 'authenticated':
        feedbacks = Feedback.objects.filter(user__isnull=False).order_by('-submitted_at')
        filter_title = "Authenticated Feedback (Lawyers Only)"
    else:
        feedbacks = Feedback.objects.all().order_by('-submitted_at')
        filter_title = "All Feedback"

    counts = {
        'all': Feedback.objects.count(),
        'lawyer': Feedback.objects.filter(user_type='lawyer').count(),
        'visitor': Feedback.objects.filter(user_type='visitor').count(),
        'anonymous': Feedback.objects.filter(user__isnull=True).count(),
        'authenticated': Feedback.objects.filter(user__isnull=False).count(),
    }

    return render(request, 'view_feedback.html', {
        'feedbacks': feedbacks,
        'selected_type': user_type,
        'filter_title': filter_title,
        'counts': counts,
    })





from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.decorators import login_required
import json

@login_required
@csrf_exempt
def clear_chat(request, user_id):
    if request.method == 'POST':
        try:
            from django.db.models import Q
            Message.objects.filter(
                Q(sender=request.user, receiver_id=user_id) |
                Q(sender_id=user_id, receiver=request.user)
            ).delete()
            
            return JsonResponse({'success': True, 'message': 'Chat cleared successfully'})
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)}, status=500)
    
    return JsonResponse({'success': False, 'error': 'Invalid request method'}, status=405)

@login_required
def delete_chat(request, user_id):
    if request.method == 'POST':
        try:
            Message.objects.filter(
                sender=request.user, 
                receiver_id=user_id
            ).update(deleted_by_sender=True)
           
            Message.objects.filter(
                sender_id=user_id, 
                receiver=request.user
            ).update(deleted_by_receiver=True)
            
            return JsonResponse({'success': True})
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)})
    
    return JsonResponse({'success': False, 'error': 'Invalid request method'})

from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
import json

@csrf_exempt
def unread_message_count(request):
    if request.user.is_authenticated:
        from .models import Message  
        count = Message.objects.filter(
            receiver=request.user,
            is_read=False
        ).count()
        return JsonResponse({'count': count})
    return JsonResponse({'count': 0})

@csrf_exempt
def mark_messages_read(request):
    if request.method == 'POST' and request.user.is_authenticated:
        from .models import Message  
        Message.objects.filter(
            receiver=request.user,
            is_read=False
        ).update(is_read=True)
        return JsonResponse({'success': True})
    return JsonResponse({'success': False})


@csrf_exempt
def get_unread_counts(request):
    if request.user.is_authenticated:
        from django.db.models import Count
        from .models import Message  
        
        
        unread_counts = Message.objects.filter(
            receiver=request.user,
            is_read=False
        ).values('sender_id').annotate(
            count=Count('id')
        )
        
        counts = {str(item['sender_id']): item['count'] for item in unread_counts}
        return JsonResponse({'counts': counts})
    return JsonResponse({'counts': {}})


@csrf_exempt
def mark_user_messages_read(request):
    if request.method == 'POST' and request.user.is_authenticated:
        try:
            data = json.loads(request.body)
            user_id = data.get('user_id')
            
            if user_id:
                from .models import Message  # Adjust import
                Message.objects.filter(
                    receiver=request.user,
                    sender_id=user_id,
                    is_read=False
                ).update(is_read=True)
                
                return JsonResponse({'success': True})
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)})
    
    return JsonResponse({'success': False})

@csrf_exempt
def get_user_unread_count(request, user_id):
    if request.user.is_authenticated:
        from .models import Message  # Adjust import
        count = Message.objects.filter(
            receiver=request.user,
            sender_id=user_id,
            is_read=False
        ).count()
        return JsonResponse({'count': count})
    return JsonResponse({'count': 0})

    
@csrf_exempt
def get_user_unread_counts(request):
    if request.user.is_authenticated and request.user.role == 'user':
        from django.db.models import Count
        from .models import Message
        
        # Get unread message counts from lawyers the user has contacted
        hired_lawyers = Hire.objects.filter(user=request.user).values_list('lawyer_id', flat=True)
        
        unread_counts = Message.objects.filter(
            receiver=request.user,
            sender_id__in=hired_lawyers,
            is_read=False
        ).values('sender_id').annotate(
            count=Count('id')
        )
        
        counts = {str(item['sender_id']): item['count'] for item in unread_counts}
        return JsonResponse({'counts': counts})
    return JsonResponse({'counts': {}})

# Enhanced chatbot views with document features
import time
import os
import json
import uuid
from datetime import datetime
from django.shortcuts import render
from django.http import JsonResponse, HttpResponse
from django.views.decorators.clickjacking import xframe_options_exempt
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.decorators import login_required
from django.conf import settings
from django.core.files.storage import default_storage
from django.core.files.base import ContentFile

# Document processing imports
import PyPDF2
import docx
from docx import Document
from docx.shared import Inches
import io

# LangChain imports
from langchain_community.vectorstores import FAISS
from langchain_google_genai import GoogleGenerativeAIEmbeddings
from langchain.prompts import PromptTemplate
from langchain_groq import ChatGroq
from langchain.memory import ConversationBufferWindowMemory
from langchain.chains import ConversationalRetrievalChain
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.docstore.document import Document as LangChainDocument

# Load environment variables
import os
from dotenv import load_dotenv
load_dotenv()

# Global variables for vector stores
try:
    embeddings = GoogleGenerativeAIEmbeddings(
        model="models/embedding-001",
        google_api_key=os.getenv("GOOGLE_API_KEY")
    )
except Exception as e:
    embeddings = None 
general_db = FAISS.load_local("my_vector_store", embeddings, allow_dangerous_deserialization=True)
document_vector_stores = {}  # Store document-specific vector stores

# Initialize the LLM
import os
groq_api_key = os.getenv("GROQ_API_KEY")
llm = ChatGroq(groq_api_key=groq_api_key, model_name="llama3-70b-8192")

# Enhanced prompt templates
general_prompt_template = """
<s>[INST]This is a chat template and As a legal chat bot, your primary objective is to provide accurate and concise information based on the user's questions. Do not generate your own questions and answers. You will adhere strictly to the instructions provided, offering relevant context from the knowledge base while avoiding unnecessary details. Your responses will be brief, to the point, and in compliance with the established format. If a question falls outside the given context, you will refrain from utilizing the chat history and instead rely on your own knowledge base to generate an appropriate response. You will prioritize the user's query and refrain from posing additional questions. The aim is to deliver professional, precise, and contextually relevant information pertaining to the Indian Penal Code.
CONTEXT: {context}
CHAT HISTORY: {chat_history}
QUESTION: {question}
ANSWER:
</s>[INST]
"""

document_prompt_template = """
<s>[INST]You are a legal document analysis assistant. Analyze the provided document context and answer questions about it accurately and concisely. Focus on the specific content of the uploaded document. If the question cannot be answered from the document content, clearly state that the information is not available in the document.

DOCUMENT CONTEXT: {context}
CHAT HISTORY: {chat_history}
QUESTION: {question}
ANSWER:
</s>[INST]
"""

generation_prompt_template = """
<s>[INST]You are a legal document generation assistant. Based on the provided information, generate a professional legal document. Follow proper legal document formatting and include all necessary clauses and provisions. Ensure the document is comprehensive and legally sound.

DOCUMENT TYPE: {document_type}
PROVIDED INFORMATION: {user_info}
GENERATE A COMPLETE LEGAL DOCUMENT:
</s>[INST]
"""

def extract_text_from_file(file):
    """Extract text from uploaded file based on file type"""
    text = ""
    file_extension = file.name.lower().split('.')[-1]
    
    try:
        if file_extension == 'pdf':
            # Read PDF
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
                
        elif file_extension in ['doc', 'docx']:
            # Read Word document
            doc = docx.Document(file)
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
                
        elif file_extension == 'txt':
            # Read text file
            text = file.read().decode('utf-8')
            
        else:
            raise ValueError(f"Unsupported file type: {file_extension}")
            
    except Exception as e:
        raise Exception(f"Error extracting text from {file.name}: {str(e)}")
    
    return text

def create_document_vector_store(text, session_key):
    """Create a vector store for the uploaded document"""
    try:
        # Split text into chunks
        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            length_function=len
        )
        chunks = text_splitter.split_text(text)
        
        # Create documents
        documents = [LangChainDocument(page_content=chunk) for chunk in chunks]
        
        # Create vector store
        vector_store = FAISS.from_documents(documents, embeddings)
        
        # Store in session-specific storage
        document_vector_stores[session_key] = vector_store
        
        return True
    except Exception as e:
        print(f"Error creating vector store: {str(e)}")
        return False

def generate_legal_document(document_type, answers):
    """Generate a legal document based on type and answers"""
    
    document_templates = {
        'contract': generate_contract,
        'lease': generate_lease_agreement,
        'nda': generate_nda,
        'power_of_attorney': generate_power_of_attorney,
        'legal_notice': generate_legal_notice,
        'will': generate_will
    }
    
    if document_type in document_templates:
        return document_templates[document_type](answers)
    else:
        raise ValueError(f"Unsupported document type: {document_type}")

def generate_contract(answers):
    """Generate a contract document"""
    doc = Document()
    
    # Title
    title = doc.add_heading('SERVICE CONTRACT', 0)
    title.alignment = 1  # Center alignment
    
    # Date and parties
    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph()
    
    doc.add_paragraph(f"This Service Contract ('Agreement') is entered into between:")
    doc.add_paragraph(f"First Party: {answers.get('party1_name', '[Party 1 Name]')}")
    doc.add_paragraph(f"Second Party: {answers.get('party2_name', '[Party 2 Name]')}")
    doc.add_paragraph()
    
    # Contract details
    doc.add_heading('1. PURPOSE', level=1)
    doc.add_paragraph(f"The purpose of this contract is: {answers.get('contract_purpose', '[Contract Purpose]')}")
    
    doc.add_heading('2. FINANCIAL TERMS', level=1)
    if answers.get('contract_amount'):
        doc.add_paragraph(f"Contract Amount: {answers.get('contract_amount')}")
    
    doc.add_heading('3. DURATION', level=1)
    doc.add_paragraph(f"Start Date: {answers.get('start_date', '[Start Date]')}")
    doc.add_paragraph(f"End Date: {answers.get('end_date', '[End Date]')}")
    
    doc.add_heading('4. TERMS AND CONDITIONS', level=1)
    doc.add_paragraph(answers.get('terms', 'Standard terms and conditions apply.'))
    
    # Signatures
    doc.add_paragraph()
    doc.add_paragraph("_" * 30 + "    " + "_" * 30)
    doc.add_paragraph("First Party Signature" + "    " * 5 + "Second Party Signature")
    
    return doc

def generate_lease_agreement(answers):
    """Generate a lease agreement document"""
    doc = Document()
    
    # Title
    title = doc.add_heading('LEASE AGREEMENT', 0)
    title.alignment = 1
    
    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph()
    
    # Parties
    doc.add_paragraph("This Lease Agreement is made between:")
    doc.add_paragraph(f"Landlord: {answers.get('landlord_name', '[Landlord Name]')}")
    doc.add_paragraph(f"Tenant: {answers.get('tenant_name', '[Tenant Name]')}")
    doc.add_paragraph()
    
    # Property details
    doc.add_heading('1. PROPERTY', level=1)
    doc.add_paragraph(f"Property Address: {answers.get('property_address', '[Property Address]')}")
    
    # Financial terms
    doc.add_heading('2. RENT AND DEPOSIT', level=1)
    doc.add_paragraph(f"Monthly Rent: ₹{answers.get('monthly_rent', '[Amount]')}")
    doc.add_paragraph(f"Security Deposit: ₹{answers.get('security_deposit', '[Amount]')}")
    
    # Lease period
    doc.add_heading('3. LEASE PERIOD', level=1)
    doc.add_paragraph(f"Lease Start Date: {answers.get('lease_start', '[Start Date]')}")
    doc.add_paragraph(f"Lease Duration: {answers.get('lease_duration', '[Duration]')} months")
    
    # Special terms
    if answers.get('special_terms'):
        doc.add_heading('4. SPECIAL TERMS', level=1)
        doc.add_paragraph(answers.get('special_terms'))
    
    # Signatures
    doc.add_paragraph()
    doc.add_paragraph("_" * 30 + "    " + "_" * 30)
    doc.add_paragraph("Landlord Signature" + "    " * 5 + "Tenant Signature")
    
    return doc

def generate_nda(answers):
    """Generate a Non-Disclosure Agreement"""
    doc = Document()
    
    title = doc.add_heading('NON-DISCLOSURE AGREEMENT', 0)
    title.alignment = 1
    
    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph()
    
    # Parties
    doc.add_paragraph("This Non-Disclosure Agreement ('Agreement') is entered into between:")
    doc.add_paragraph(f"Disclosing Party: {answers.get('disclosing_party', '[Disclosing Party]')}")
    doc.add_paragraph(f"Receiving Party: {answers.get('receiving_party', '[Receiving Party]')}")
    doc.add_paragraph()
    
    # Confidential information
    doc.add_heading('1. CONFIDENTIAL INFORMATION', level=1)
    doc.add_paragraph(f"Confidential Information includes: {answers.get('confidential_info', '[Description]')}")
    
    # Purpose
    doc.add_heading('2. PURPOSE', level=1)
    doc.add_paragraph(f"Purpose of disclosure: {answers.get('purpose', '[Purpose]')}")
    
    # Duration
    doc.add_heading('3. DURATION', level=1)
    doc.add_paragraph(f"This agreement shall remain in effect for {answers.get('duration', '[Duration]')} years.")
    
    # Governing law
    doc.add_heading('4. GOVERNING LAW', level=1)
    doc.add_paragraph(f"This agreement shall be governed by the laws of {answers.get('jurisdiction', '[Jurisdiction]')}.")
    
    # Signatures
    doc.add_paragraph()
    doc.add_paragraph("_" * 30 + "    " + "_" * 30)
    doc.add_paragraph("Disclosing Party" + "    " * 5 + "Receiving Party")
    
    return doc

def generate_power_of_attorney(answers):
    """Generate a Power of Attorney document"""
    doc = Document()
    
    title = doc.add_heading('POWER OF ATTORNEY', 0)
    title.alignment = 1
    
    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph()
    
    # Parties
    doc.add_paragraph("KNOW ALL BY THESE PRESENTS:")
    doc.add_paragraph(f"I, {answers.get('principal_name', '[Principal Name]')}, of sound mind and legal age, hereby appoint {answers.get('agent_name', '[Agent Name]')} as my Attorney-in-Fact.")
    doc.add_paragraph()
    
    # Powers granted
    doc.add_heading('POWERS GRANTED', level=1)
    doc.add_paragraph(f"The following powers are hereby granted: {answers.get('powers_granted', '[Powers Description]')}")
    
    # Effective date
    doc.add_heading('EFFECTIVE DATE', level=1)
    doc.add_paragraph(f"This Power of Attorney shall be effective from: {answers.get('effective_date', '[Date]')}")
    
    # Duration
    doc.add_heading('DURATION', level=1)
    duration_type = answers.get('duration_type', 'Limited')
    if duration_type == 'Durable':
        doc.add_paragraph("This is a DURABLE Power of Attorney and shall remain in effect even if I become incapacitated.")
    else:
        doc.add_paragraph("This is a LIMITED Power of Attorney.")
    
    # Limitations
    if answers.get('limitations'):
        doc.add_heading('LIMITATIONS', level=1)
        doc.add_paragraph(answers.get('limitations'))
    
    # Signature
    doc.add_paragraph()
    doc.add_paragraph("_" * 40)
    doc.add_paragraph(f"Principal: {answers.get('principal_name', '[Principal Name]')}")
    
    return doc

def generate_legal_notice(answers):
    """Generate a Legal Notice"""
    doc = Document()
    
    title = doc.add_heading('LEGAL NOTICE', 0)
    title.alignment = 1
    
    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph()
    
    # To and From
    doc.add_paragraph(f"TO: {answers.get('recipient_name', '[Recipient Name]')}")
    doc.add_paragraph(f"FROM: {answers.get('sender_name', '[Sender Name]')}")
    doc.add_paragraph()
    
    # Subject
    doc.add_heading('SUBJECT', level=1)
    doc.add_paragraph(answers.get('notice_subject', '[Subject]'))
    
    # Details
    doc.add_heading('NOTICE', level=1)
    doc.add_paragraph("You are hereby notified that:")
    doc.add_paragraph(answers.get('violation_details', '[Details of the issue]'))
    
    # Remedy sought
    doc.add_heading('REMEDY SOUGHT', level=1)
    doc.add_paragraph(answers.get('remedy_sought', '[Remedy description]'))
    
    # Deadline
    doc.add_heading('RESPONSE DEADLINE', level=1)
    doc.add_paragraph(f"You are required to respond to this notice by {answers.get('deadline', '[Date]')}.")
    doc.add_paragraph("Failure to respond may result in legal action being taken against you.")
    
    # Signature
    doc.add_paragraph()
    doc.add_paragraph("_" * 30)
    doc.add_paragraph(f"Sent by: {answers.get('sender_name', '[Sender Name]')}")
    
    return doc

def generate_will(answers):
    """Generate a Will/Testament"""
    doc = Document()
    
    title = doc.add_heading('LAST WILL AND TESTAMENT', 0)
    title.alignment = 1
    
    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph()
    
    # Testator
    doc.add_paragraph(f"I, {answers.get('testator_name', '[Testator Name]')}, being of sound mind and memory, do hereby make, publish, and declare this to be my Last Will and Testament.")
    doc.add_paragraph()
    
    # Executor
    doc.add_heading('EXECUTOR', level=1)
    doc.add_paragraph(f"I hereby nominate and appoint {answers.get('executor_name', '[Executor Name]')} as the Executor of this Will.")
    
    # Beneficiaries
    doc.add_heading('BENEFICIARIES AND DISTRIBUTION', level=1)
    doc.add_paragraph("I give, devise, and bequeath my estate as follows:")
    doc.add_paragraph(answers.get('beneficiaries', '[Beneficiary details]'))
    
    # Assets
    doc.add_heading('ASSETS', level=1)
    doc.add_paragraph(f"Description of assets: {answers.get('assets_description', '[Asset description]')}")
    
    # Special bequests
    if answers.get('special_bequests'):
        doc.add_heading('SPECIAL BEQUESTS', level=1)
        doc.add_paragraph(answers.get('special_bequests'))
    
    # Guardian for minors
    if answers.get('guardian_minor'):
        doc.add_heading('GUARDIAN FOR MINOR CHILDREN', level=1)
        doc.add_paragraph(f"I appoint {answers.get('guardian_minor')} as guardian for any minor children.")
    
    # Signature
    doc.add_paragraph()
    doc.add_paragraph("IN WITNESS WHEREOF, I have executed this Will.")
    doc.add_paragraph()
    doc.add_paragraph("_" * 40)
    doc.add_paragraph(f"Testator: {answers.get('testator_name', '[Testator Name]')}")
    doc.add_paragraph()
    doc.add_paragraph("WITNESSES:")
    doc.add_paragraph("_" * 30 + "    " + "_" * 30)
    doc.add_paragraph("Witness 1" + "    " * 8 + "Witness 2")
    
    return doc

@xframe_options_exempt
@csrf_exempt
def chatbot(request):
    if request.method == 'POST':
        # Handle different actions
        action = request.POST.get('action')
        
        if action == 'upload_document':
            return handle_document_upload(request)
        elif action == 'generate_document':
            return handle_document_generation(request)
        else:
            return handle_chat_message(request)
    
    return render(request, 'chat1.html')

def handle_document_upload(request):
    """Handle document upload and processing"""
    try:
        document = request.FILES.get('document')
        if not document:
            return JsonResponse({'success': False, 'error': 'No document provided'})
        
        # Extract text from document
        text = extract_text_from_file(document)
        
        if not text.strip():
            return JsonResponse({'success': False, 'error': 'Could not extract text from document'})
        
        # Create vector store for this session
        session_key = request.session.session_key or request.session.create()
        
        if create_document_vector_store(text, session_key):
            return JsonResponse({'success': True, 'message': 'Document processed successfully'})
        else:
            return JsonResponse({'success': False, 'error': 'Failed to process document'})
            
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})

def handle_document_generation(request):
    """Handle legal document generation"""
    try:
        document_type = request.POST.get('document_type')
        answers_json = request.POST.get('answers', '{}')
        answers = json.loads(answers_json)
        
        # Generate the document
        doc = generate_legal_document(document_type, answers)
        
        # Save document to temporary storage
        document_id = str(uuid.uuid4())
        temp_path = os.path.join(settings.MEDIA_ROOT, 'temp_documents', f'{document_id}.docx')
        
        # Ensure directory exists
        os.makedirs(os.path.dirname(temp_path), exist_ok=True)
        
        # Save document
        doc.save(temp_path)
        
        # Store document info in session
        if 'generated_documents' not in request.session:
            request.session['generated_documents'] = {}
        
        request.session['generated_documents'][document_id] = {
            'path': temp_path,
            'type': document_type,
            'created_at': datetime.now().isoformat()
        }
        request.session.modified = True
        
        return JsonResponse({
            'success': True,
            'document_id': document_id,
            'message': 'Document generated successfully'
        })
        
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})

def handle_chat_message(request):
    """Handle regular chat messages"""
    try:
        user_input = request.POST.get('user_input')
        mode = request.POST.get('mode', 'general')
        has_document = request.POST.get('has_document') == 'true'
        
        # Get or create session
        session_key = request.session.session_key or request.session.create()
        
        # Initialize chat history
        if 'chat_history' not in request.session:
            request.session['chat_history'] = []
        
        chat_history = request.session['chat_history']
        
        # Choose appropriate chain based on mode
        if mode == 'document' and has_document and session_key in document_vector_stores:
            # Document Q&A mode
            document_db = document_vector_stores[session_key]
            retriever = document_db.as_retriever(search_type="similarity", search_kwargs={"k": 3})
            
            prompt = PromptTemplate(
                template=document_prompt_template,
                input_variables=['context', 'question', 'chat_history']
            )
            
            memory = ConversationBufferWindowMemory(
                k=3, 
                memory_key="chat_history", 
                return_messages=True
            )
            
            qa_chain = ConversationalRetrievalChain.from_llm(
                llm=llm,
                memory=memory,
                retriever=retriever,
                combine_docs_chain_kwargs={'prompt': prompt}
            )
            
        else:
            # General legal chat mode
            general_retriever = general_db.as_retriever(search_type="similarity", search_kwargs={"k": 4})
            
            prompt = PromptTemplate(
                template=general_prompt_template,
                input_variables=['context', 'question', 'chat_history']
            )
            
            memory = ConversationBufferWindowMemory(
                k=2, 
                memory_key="chat_history", 
                return_messages=True
            )
            
            qa_chain = ConversationalRetrievalChain.from_llm(
                llm=llm,
                memory=memory,
                retriever=general_retriever,
                combine_docs_chain_kwargs={'prompt': prompt}
            )
        
        # Get response
        result = qa_chain.invoke(input=user_input)
        response_text = result["answer"]
        
        # Update chat history
        chat_history.append({"role": "user", "content": user_input})
        chat_history.append({"role": "assistant", "content": response_text})
        
        # Keep only last 10 messages
        if len(chat_history) > 10:
            chat_history = chat_history[-10:]
        
        request.session['chat_history'] = chat_history
        request.session.modified = True
        
        return JsonResponse({'answer': response_text})
        
    except Exception as e:
        return JsonResponse({'answer': f'Sorry, I encountered an error: {str(e)}'})

@login_required
def download_document(request, document_id):
    """Download generated document"""
    try:
        if 'generated_documents' not in request.session:
            return HttpResponse('Document not found', status=404)
        
        doc_info = request.session['generated_documents'].get(document_id)
        if not doc_info:
            return HttpResponse('Document not found', status=404)
        
        file_path = doc_info['path']
        if not os.path.exists(file_path):
            return HttpResponse('Document file not found', status=404)
        
        # Read file content
        with open(file_path, 'rb') as f:
            response = HttpResponse(
                f.read(),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )
            response['Content-Disposition'] = f'attachment; filename="{doc_info["type"]}_{document_id}.docx"'
            return response
            
    except Exception as e:
        return HttpResponse(f'Error downloading document: {str(e)}', status=500)
    




import random
from django.core.mail import send_mail
from django.contrib import messages

def forgot_password(request):
    if request.method == 'POST':
        email = request.POST['email']
        user = CustomUser.objects.filter(email=email).first()

        if user:
            # Generate a 6-digit OTP
            otp = ''.join(random.choices('0123456789', k=6))

            # Save the OTP in the user object
            user.otp = otp
            user.save()

            # Send the OTP to the user's email
            subject = 'Password Reset OTP'
            message = f'Your OTP is: {otp}'
            from_email = 'ramenterprise9a@gmail.com'
            recipient_list = [email]
            send_mail(subject, message, from_email, recipient_list)

            # Render the OTP verification form
            return render(request, 'verify_otp.html', {'email': email})
        else:
            message = "This email is not registered!"
            return render(request, 'forgot_password.html', {'msg': message})

    return render(request, 'forgot_password.html')

def verify_otp(request):
    if request.method == 'POST':
        email = request.POST['email']
        otp = request.POST['otp']
        user = CustomUser.objects.filter(email=email, otp=otp).first()

        if user:
            # Clear the OTP field
            user.otp = None
            user.save()

            # Render the password reset form
            return render(request, 'reset_password.html', {'email': email})
        else:
            message = "Invalid OTP!"
            return render(request, 'verify_otp.html', {'email': email, 'msg': message})

    return render(request, 'verify_otp.html')

def update_password(request):
    if request.method == 'POST':
        email = request.POST['email']
        new_password = request.POST['new_password']
        confirm_password = request.POST['confirm_password']

        if new_password == confirm_password:
            # Find the user by email
            user = CustomUser.objects.filter(email=email).first()

            if user:
                # Use Django's set_password method to properly hash the password
                user.set_password(new_password)
                user.save()

                # Add success message
                messages.success(request, "Password updated successfully! Please login with your new password.")
                return render(request, 'login.html')
            else:
                message = "User not found!"
                return render(request, 'reset_password.html', {'email': email, 'msg': message})
        else:
            message = "Passwords do not match!"
            return render(request, 'reset_password.html', {'email': email, 'msg': message})

    return render(request, 'reset_password.html')


